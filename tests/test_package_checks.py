"""Checks that need more than one worksheet, or the csv files.

`ExcelReader` is built per worksheet and by design sees nothing else; these
two live in the processor, which holds every question list and knows where
the csv files are.
"""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from crf_reader import CRFS_COLUMN_NAMES
from models import AppConfig

from processor import SurveyGenProcessor
from tests.test_dd_validations import HEADERS, row
from tests.test_processor_gating import run_quietly

VILLAGES = "region,vcode,villagename\n1,11,Alpha\n1,12,Beta\n2,21,Gamma\n"


def csv_question(fieldname, filter_column="region", value="vcode", display="villagename"):
    responses = (
        "source: csv\nfile: villages.csv\n"
        f"filter: {filter_column} = [[region]]\ndisplay: {display}\nvalue: {value}"
    )
    return row(fieldname, "radio", "integer", responses=responses)


def build(tmp, sheets, csv_text=VILLAGES):
    """`sheets` is {worksheet title: rows}; every sheet gets a crfs row."""
    tmp = Path(tmp)
    workbook = Workbook()
    first = True
    crfs_rows = []
    for order, (title, rows) in enumerate(sheets.items(), start=1):
        ws = workbook.active if first else workbook.create_sheet(title)
        ws.title = title
        first = False
        ws.append(HEADERS)
        ws.append(row("code", "automatic", "text", "Code"))
        for r in rows:
            ws.append(r)
        table = title.replace("_dd", "")
        crfs_rows.append([order * 10, table, table, "code",
                          '{"prefix": "C", "fields": [], "incrementLength": 4}',
                          1 if order == 1 else 0, "", "", "", "", "", "", "", "code", ""])
    crfs = workbook.create_sheet("crfs")
    crfs.append(list(CRFS_COLUMN_NAMES))
    for r in crfs_rows:
        crfs.append(r)
    excel = tmp / "dictionary.xlsx"
    workbook.save(excel)
    workbook.close()
    csv_dir = tmp / "csv"
    csv_dir.mkdir()
    (csv_dir / "villages.csv").write_text(csv_text, encoding="utf-8")
    processor = SurveyGenProcessor(AppConfig(
        excelFile=str(excel), csvFiles=str(csv_dir), outputPath=str(tmp / "out"),
        surveyName="Test", surveyId="test_survey", databaseName="test.sqlite",
    ))
    code = run_quietly(processor)
    return code, processor.logstring


def errors(log):
    return [line for line in log if line.startswith("ERROR")]


def warnings(log):
    return [line for line in log if line.startswith("WARNING")]


class CsvColumnTests(unittest.TestCase):
    def test_a_filter_on_a_missing_column_is_an_error(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {"form_dd": [
                row("region", "radio", "integer", responses="1:One\n2:Two"),
                csv_question("village", filter_column="regoin"),
            ]})
        self.assertEqual(code, 1)
        message = "\n".join(errors(log))
        self.assertIn("regoin", message)
        self.assertIn("always empty", message)

    def test_a_real_column_is_silent(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {"form_dd": [
                row("region", "radio", "integer", responses="1:One\n2:Two"),
                csv_question("village"),
            ]})
        self.assertEqual(code, 0, "\n".join(log))


class CsvSkipValueTests(unittest.TestCase):
    def test_a_skip_on_a_value_the_csv_never_holds_is_an_error(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {"form_dd": [
                row("region", "radio", "integer", responses="1:One\n2:Two"),
                csv_question("village"),
                row("place", "text", "text", maxchars="20",
                    skip="preskip: if village = 999, skip to after"),
                row("after", "text", "text", maxchars="20"),
            ]})
        self.assertEqual(code, 1)
        message = "\n".join(errors(log))
        self.assertIn("against 999", message)
        self.assertIn("villages.csv", message)
        self.assertIn("stay silent", message)

    def test_a_declared_not_in_list_code_is_a_value(self):
        with TemporaryDirectory() as tmp:
            q = csv_question("village")
            q[5] += "\nnot_in_list: 999, Not listed"
            code, log = build(tmp, {"form_dd": [
                row("region", "radio", "integer", responses="1:One\n2:Two"),
                q,
                row("place", "text", "text", maxchars="20",
                    skip="preskip: if village = 999, skip to after"),
                row("after", "text", "text", maxchars="20"),
            ]})
        self.assertEqual(code, 0, "\n".join(log))

    def test_a_value_in_the_csv_is_silent(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {"form_dd": [
                row("region", "radio", "integer", responses="1:One\n2:Two"),
                csv_question("village"),
                row("place", "text", "text", maxchars="20",
                    skip="preskip: if village = 21, skip to after"),
                row("after", "text", "text", maxchars="20"),
            ]})
        self.assertEqual(code, 0, "\n".join(log))


class FieldsAcrossFormsTests(unittest.TestCase):
    def test_the_same_field_with_different_codes_warns(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {
                "first_dd": [row("sex", "radio", "integer", responses="1:M\n2:F")],
                "second_dd": [row("sex", "radio", "integer", responses="1:M\n2:F\n9:Unknown")],
            })
        self.assertEqual(code, 0, "\n".join(log))
        message = "\n".join(warnings(log))
        self.assertIn("'sex' is defined differently", message)
        self.assertIn("first_dd", message)
        self.assertIn("second_dd", message)

    def test_the_same_definition_twice_is_silent(self):
        with TemporaryDirectory() as tmp:
            code, log = build(tmp, {
                "first_dd": [row("sex", "radio", "integer", responses="1:M\n2:F")],
                "second_dd": [row("sex", "radio", "integer", responses="1:M\n2:F")],
            })
        self.assertEqual(code, 0, "\n".join(log))
        self.assertEqual([w for w in warnings(log) if "defined differently" in w], [])
